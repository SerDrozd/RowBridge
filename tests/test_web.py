from __future__ import annotations

import re
from pathlib import Path

from fastapi.testclient import TestClient

from rowbridge import __version__
from rowbridge.config import Settings
from rowbridge.web import create_app

A_CSV = (
    b"invoice_ref,customer,amount,date\n"
    b"INV-1,Acme Ltd,10.00,2026-10-01\n"
    b"INV-2,Blue Finch GmbH,20.00,2026-10-02\n"
)
B_CSV = (
    b"reference,payer,total,paid_at\n"
    b"INV1,ACME Limited,10.01,2026-10-02\n"
    b"INV-9,Blue Finch GmbH,20.00,2026-10-02\n"
)


def make_client(tmp_path: Path) -> TestClient:
    app = create_app(Settings(data_dir=tmp_path / "data"))
    return TestClient(app)


def create_sample_run(client: TestClient) -> str:
    prepared = client.post(
        "/prepare",
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": ("payments.csv", B_CSV, "text/csv"),
        },
    )
    assert prepared.status_code == 200
    assert "Secondary text" in prepared.text
    stage_match = re.search(r'data-stage-id="([a-f0-9]+)"', prepared.text)
    assert stage_match is not None

    created = client.post(
        "/runs",
        data={
            "stage_id": stage_match.group(1),
            "primary_a": "invoice_ref",
            "primary_b": "reference",
            "secondary_a": "customer",
            "secondary_b": "payer",
            "amount_a": "amount",
            "amount_b": "total",
            "date_a": "date",
            "date_b": "paid_at",
            "amount_tolerance": "0.05",
            "date_window_days": "2",
        },
        follow_redirects=False,
    )
    assert created.status_code == 303
    return created.headers["location"]


def review_match_id(html: str) -> int:
    match = re.search(r'/matches/(\d+)/accept', html)
    assert match is not None
    return int(match.group(1))


def test_runs_page_shows_empty_state_and_header_link(tmp_path: Path) -> None:
    client = make_client(tmp_path)

    home = client.get("/")
    history = client.get("/runs")

    assert home.status_code == 200
    assert 'href="/runs">Runs</a>' in home.text
    assert history.status_code == 200
    assert "Reconciliation runs" in history.text
    assert "No saved runs" in history.text
    assert "Start a reconciliation" in history.text

    missing = client.get("/runs?page=2")
    assert missing.status_code == 404
    assert "Run history page not found" in missing.text


def test_runs_page_lists_saved_run_with_current_summary(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)
    accepted = client.post(
        f"{location}/matches/{match_id}/accept",
        follow_redirects=False,
    )
    assert accepted.status_code == 303

    history = client.get("/runs")

    assert history.status_code == 200
    assert "orders.csv" in history.text
    assert "payments.csv" in history.text
    assert f'href="{location}">Open</a>' in history.text
    assert '<span class="run-count run-count-good">1</span>' in history.text
    assert '<span class="run-count run-count-review">0</span>' in history.text


def test_runs_page_paginates_and_rejects_out_of_range_pages(tmp_path: Path) -> None:
    from datetime import UTC, datetime, timedelta

    from rowbridge.models import FieldMapping, MatchDecision, MatchSettings, MatchStatus
    from rowbridge.storage import Repository

    settings = Settings(data_dir=tmp_path / "data", runs_page_size=2)
    client = TestClient(create_app(settings))
    repository = Repository(settings.database_path)

    for index in range(3):
        repository.save_run(
            run_id=str(index) * 32,
            created_at=(datetime(2026, 10, 4, 9, tzinfo=UTC) + timedelta(hours=index)).isoformat(
                timespec="seconds"
            ),
            filename_a=f"a-{index}.csv",
            filename_b=f"b-{index}.csv",
            mapping=FieldMapping(primary_a="id", primary_b="id"),
            settings=MatchSettings(),
            rows_a=({"id": f"A-{index}"},),
            rows_b=({"id": f"B-{index}"},),
            decisions=(MatchDecision(0, 0, 1.0, MatchStatus.AUTO_MATCHED, ()),),
        )

    first = client.get("/runs")
    second = client.get("/runs?page=2")
    missing = client.get("/runs?page=3")

    assert first.status_code == 200
    assert "a-2.csv" in first.text
    assert "a-1.csv" in first.text
    assert "a-0.csv" not in first.text
    assert 'href="/runs?page=2">Next</a>' in first.text

    assert second.status_code == 200
    assert "a-0.csv" in second.text
    assert "a-2.csv" not in second.text
    assert 'href="/runs?page=1">Previous</a>' in second.text

    assert missing.status_code == 404
    assert "Run history page not found" in missing.text


def test_delete_run_confirmation_shows_run_and_cancel_link(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)

    history = client.get("/runs")
    confirmation = client.get(f"{location}/delete")

    assert history.status_code == 200
    assert f'href="{location}/delete">Delete</a>' in history.text
    assert confirmation.status_code == 200
    assert "Delete this reconciliation run?" in confirmation.text
    assert "orders.csv" in confirmation.text
    assert "payments.csv" in confirmation.text
    assert f'href="{location}">Cancel</a>' in confirmation.text
    assert f'action="{location}/delete"' in confirmation.text


def test_delete_run_post_removes_run_and_redirects_to_history(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    deleted_location = create_sample_run(client)
    kept_location = create_sample_run(client)

    deleted = client.post(f"{deleted_location}/delete", follow_redirects=False)

    assert deleted.status_code == 303
    assert deleted.headers["location"] == "/runs"
    assert client.get(deleted_location).status_code == 404
    assert client.get(kept_location).status_code == 200

    history = client.get("/runs")
    assert history.status_code == 200
    assert history.text.count("orders.csv") == 1
    assert history.text.count("payments.csv") == 1


def test_delete_missing_run_returns_not_found(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    missing = "f" * 32

    confirmation = client.get(f"/runs/{missing}/delete")
    deleted = client.post(f"/runs/{missing}/delete", follow_redirects=False)

    assert confirmation.status_code == 404
    assert deleted.status_code == 404
    assert "Run not found" in confirmation.text
    assert "Run not found" in deleted.text


def test_vertical_slice_upload_map_persist_render_and_export(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)

    results = client.get(location)
    assert results.status_code == 200
    assert "Reconciliation results" in results.text
    assert "INV-1" in results.text
    assert "INV1" in results.text
    assert "Blue Finch GmbH" in results.text
    assert ">Accept<" in results.text
    assert "Link unmatched rows" in results.text

    fresh_client = make_client(tmp_path)
    persisted = fresh_client.get(location)
    assert persisted.status_code == 200
    assert "INV-1" in persisted.text

    export = fresh_client.get(f"{location}/export.csv")
    assert export.status_code == 200
    assert "status,score,side_a_row" in export.text
    assert "INV-1,INV1" in export.text


def test_accept_review_persists_confirmation_and_audit_event(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)

    accepted = client.post(
        f"{location}/matches/{match_id}/accept",
        follow_redirects=False,
    )
    assert accepted.status_code == 303

    fresh_client = make_client(tmp_path)
    persisted = fresh_client.get(location)
    assert persisted.status_code == 200
    assert "confirmed" in persisted.text
    assert "Accepted proposed match." in persisted.text
    assert "accept" in persisted.text
    assert "<strong>0</strong><span>Needs review</span>" in persisted.text
    assert "<strong>1</strong><span>Human matched</span>" in persisted.text

    repeated = fresh_client.post(
        f"{location}/matches/{match_id}/accept",
        follow_redirects=False,
    )
    assert repeated.status_code == 409


def test_confirmed_match_can_be_reopened_from_results(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)

    accepted = client.post(
        f"{location}/matches/{match_id}/accept",
        follow_redirects=False,
    )
    assert accepted.status_code == 303

    confirmed = client.get(location)
    assert f'action="{location}/matches/{match_id}/reopen"' in confirmed.text
    assert ">Reopen review<" in confirmed.text

    reopened = client.post(
        f"{location}/matches/{match_id}/reopen",
        follow_redirects=False,
    )
    assert reopened.status_code == 303
    assert reopened.headers["location"] == f"{location}#review-history"

    persisted = client.get(location)
    assert f'action="{location}/matches/{match_id}/accept"' in persisted.text
    assert "Reopened confirmed match for review." in persisted.text
    assert "<strong>1</strong><span>Needs review</span>" in persisted.text
    assert "<strong>0</strong><span>Human matched</span>" in persisted.text


def test_rejected_proposal_can_be_restored_from_review_history(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)

    rejected = client.post(
        f"{location}/matches/{match_id}/reject",
        follow_redirects=False,
    )
    assert rejected.status_code == 303

    after_reject = client.get(location)
    restore_action = re.search(
        rf'action="{re.escape(location)}/review-events/(\d+)/restore"',
        after_reject.text,
    )
    assert restore_action is not None
    assert ">Restore proposal<" in after_reject.text

    restored = client.post(
        f"{location}/review-events/{restore_action.group(1)}/restore",
        follow_redirects=False,
    )
    assert restored.status_code == 303
    assert restored.headers["location"] == f"{location}?view=review#review-history"

    persisted = client.get(location)
    assert "Restored rejected proposal to review." in persisted.text
    assert "<strong>1</strong><span>Needs review</span>" in persisted.text
    assert "<strong>0</strong><span>Unmatched rows</span>" in persisted.text
    assert ">Restore proposal<" not in persisted.text


def test_manual_match_can_be_unlinked_from_results(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)

    rejected = client.post(
        f"{location}/matches/{match_id}/reject",
        follow_redirects=False,
    )
    assert rejected.status_code == 303

    after_reject = client.get(location)
    a_option = re.search(r'<option value="(\d+)">row 3 · INV-2</option>', after_reject.text)
    b_option = re.search(r'<option value="(\d+)">row 3 · INV-9</option>', after_reject.text)
    assert a_option is not None
    assert b_option is not None

    linked = client.post(
        f"{location}/manual-links",
        data={
            "a_match_id": a_option.group(1),
            "b_match_id": b_option.group(1),
        },
        follow_redirects=False,
    )
    assert linked.status_code == 303

    manual = client.get(location)
    unlink_action = re.search(
        rf'action="{re.escape(location)}/matches/(\d+)/unlink"',
        manual.text,
    )
    assert unlink_action is not None
    assert ">Unlink<" in manual.text

    unlinked = client.post(
        f"{location}/matches/{unlink_action.group(1)}/unlink",
        follow_redirects=False,
    )
    assert unlinked.status_code == 303
    assert unlinked.headers["location"] == f"{location}#manual-link"

    persisted = client.get(location)
    assert "Unlinked manual match; both rows returned to unmatched." in persisted.text
    assert "<strong>0</strong><span>Human matched</span>" in persisted.text
    assert "<strong>2</strong><span>Unmatched rows</span>" in persisted.text
    assert "row 3 · INV-2" in persisted.text
    assert "row 3 · INV-9" in persisted.text


def test_reversible_review_routes_reject_invalid_transitions(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)

    reopen_review = client.post(
        f"{location}/matches/{match_id}/reopen",
        follow_redirects=False,
    )
    unlink_review = client.post(
        f"{location}/matches/{match_id}/unlink",
        follow_redirects=False,
    )
    missing_restore = client.post(
        f"{location}/review-events/999999/restore",
        follow_redirects=False,
    )

    assert reopen_review.status_code == 409
    assert unlink_review.status_code == 409
    assert missing_restore.status_code == 409


def test_reject_review_then_manually_link_unmatched_rows(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)

    rejected = client.post(
        f"{location}/matches/{match_id}/reject",
        follow_redirects=False,
    )
    assert rejected.status_code == 303

    after_reject = client.get(location)
    assert "Rejected proposed match" in after_reject.text
    assert "<strong>0</strong><span>Needs review</span>" in after_reject.text
    assert "<strong>2</strong><span>Unmatched rows</span>" in after_reject.text

    a_option = re.search(r'<option value="(\d+)">row 3 · INV-2</option>', after_reject.text)
    b_option = re.search(r'<option value="(\d+)">row 3 · INV-9</option>', after_reject.text)
    assert a_option is not None
    assert b_option is not None

    linked = client.post(
        f"{location}/manual-links",
        data={
            "a_match_id": a_option.group(1),
            "b_match_id": b_option.group(1),
        },
        follow_redirects=False,
    )
    assert linked.status_code == 303

    fresh_client = make_client(tmp_path)
    persisted = fresh_client.get(location)
    assert "manual matched" in persisted.text
    assert "Linked two previously unmatched rows manually." in persisted.text
    assert "<strong>1</strong><span>Human matched</span>" in persisted.text
    assert "<strong>0</strong><span>Unmatched rows</span>" in persisted.text

    export = fresh_client.get(f"{location}/export.csv")
    assert "manual_matched,,3,3,INV-2,INV-9" in export.text


def test_result_filters_reject_unknown_view(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)

    review = client.get(f"{location}?view=review")
    assert review.status_code == 200
    assert "INV-2" in review.text
    assert "INV-1" not in review.text

    invalid = client.get(f"{location}?view=wat")
    assert invalid.status_code == 400



def test_home_file_picker_only_uses_explicit_browse_labels(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.get("/")
    assert response.status_code == 200
    assert 'id="rb-file-a"' in response.text
    assert 'data-file-trigger="rb-file-a"' in response.text
    assert 'id="rb-file-b"' in response.text
    assert 'data-file-trigger="rb-file-b"' in response.text
    assert 'hidden id="rb-file-a"' in response.text
    assert 'hidden id="rb-file-b"' in response.text
    assert f"app.css?v={__version__}" in response.text
    assert f"app.js?v={__version__}" in response.text
    assert (
        response.text.count(
            "CSV or XLSX · up to 5 MB · up to 50,000 rows · up to 200 columns"
        )
        == 2
    )
    assert 'class="file-input"' not in response.text

def test_prepare_rejects_unsupported_files(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.post(
        "/prepare",
        files={
            "file_a": ("a.txt", A_CSV, "text/plain"),
            "file_b": ("b.csv", B_CSV, "text/csv"),
        },
    )
    assert response.status_code == 400


def xlsx_bytes(headers: list[str], rows: list[list[object]]) -> bytes:
    import io

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_prepare_accepts_mixed_csv_and_xlsx_and_shows_preview(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    side_b = xlsx_bytes(
        ["reference", "payer", "total", "paid_at"],
        [["INV1", "ACME Limited", 10.01, "2026-10-02"]],
    )

    response = client.post(
        "/prepare",
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": (
                "payments.xlsx",
                side_b,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )

    assert response.status_code == 200
    assert "Map the fields that mean the same thing" in response.text
    assert "CSV · utf-8 · comma delimiter" in response.text
    assert "Excel workbook · sheet Data" in response.text
    assert "ACME Limited" in response.text


def test_export_workbook_contains_reconciliation_and_audit_sheets(tmp_path: Path) -> None:
    import io

    from openpyxl import load_workbook

    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)
    accepted = client.post(f"{location}/matches/{match_id}/accept", follow_redirects=False)
    assert accepted.status_code == 303

    exported = client.get(f"{location}/export.xlsx")
    assert exported.status_code == 200
    assert "spreadsheetml.sheet" in exported.headers["content-type"]

    workbook = load_workbook(io.BytesIO(exported.content), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [
            "Summary",
            "Reconciliation",
            "Unmatched A",
            "Unmatched B",
            "Review history",
        ]
        reconciliation = workbook["Reconciliation"]
        headers = [cell.value for cell in next(reconciliation.iter_rows())]
        assert headers[:6] == [
            "status",
            "score",
            "side_a_row",
            "side_b_row",
            "side_a_primary",
            "side_b_primary",
        ]
        review = workbook["Review history"]
        review_values = list(review.values)
        assert any(row[1] == "accept" for row in review_values[1:])
    finally:
        workbook.close()


def test_prepare_rejects_unsupported_file_type_with_clear_error(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.post(
        "/prepare",
        files={
            "file_a": ("a.txt", A_CSV, "text/plain"),
            "file_b": ("b.csv", B_CSV, "text/csv"),
        },
    )
    assert response.status_code == 400
    assert "Use a .csv or .xlsx file" in response.text


def test_responses_include_local_safety_headers(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.get("/")

    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "DENY"
    assert response.headers["referrer-policy"] == "no-referrer"
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]


def test_local_post_allows_unusual_origin_when_fetch_metadata_is_same_origin(
    tmp_path: Path,
) -> None:
    app = create_app(Settings(data_dir=tmp_path / "data"))
    client = TestClient(app, base_url="http://127.0.0.1:8000")
    response = client.post(
        "/prepare",
        headers={
            "origin": "null",
            "sec-fetch-site": "same-origin",
        },
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": ("payments.csv", B_CSV, "text/csv"),
        },
    )

    assert response.status_code == 200
    assert "Map the fields that mean the same thing" in response.text


def test_local_post_without_fetch_metadata_remains_compatible(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.post(
        "/prepare",
        headers={"origin": "null"},
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": ("payments.csv", B_CSV, "text/csv"),
        },
    )

    assert response.status_code == 200


def test_cross_site_post_is_rejected_by_fetch_metadata(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.post(
        "/prepare",
        headers={
            "origin": "https://example.com",
            "sec-fetch-site": "cross-site",
        },
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": ("payments.csv", B_CSV, "text/csv"),
        },
    )

    assert response.status_code == 403
    assert "Cross-site requests are not allowed" in response.text


def test_oversized_upload_is_rejected_before_parsing(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    too_large = b"a" * (5 * 1024 * 1024 + 1)
    response = client.post(
        "/prepare",
        files={
            "file_a": ("orders.csv", too_large, "text/csv"),
            "file_b": ("payments.csv", B_CSV, "text/csv"),
        },
    )

    assert response.status_code == 400
    assert "larger than 5 MB" in response.text


def test_successful_run_removes_staged_upload(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    prepared = client.post(
        "/prepare",
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": ("payments.csv", B_CSV, "text/csv"),
        },
    )
    stage_match = re.search(r'data-stage-id="([a-f0-9]+)"', prepared.text)
    assert stage_match is not None
    stage_id = stage_match.group(1)
    stage_dir = tmp_path / "data" / "uploads" / stage_id
    assert stage_dir.is_dir()

    created = client.post(
        "/runs",
        data={
            "stage_id": stage_id,
            "primary_a": "invoice_ref",
            "primary_b": "reference",
            "secondary_a": "customer",
            "secondary_b": "payer",
            "amount_a": "amount",
            "amount_b": "total",
            "date_a": "date",
            "date_b": "paid_at",
            "amount_tolerance": "0.05",
            "date_window_days": "2",
        },
        follow_redirects=False,
    )

    assert created.status_code == 303
    assert not stage_dir.exists()


def test_duplicate_mapping_roles_return_clear_error(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    prepared = client.post(
        "/prepare",
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": ("payments.csv", B_CSV, "text/csv"),
        },
    )
    stage_match = re.search(r'data-stage-id="([a-f0-9]+)"', prepared.text)
    assert stage_match is not None

    response = client.post(
        "/runs",
        data={
            "stage_id": stage_match.group(1),
            "primary_a": "invoice_ref",
            "primary_b": "reference",
            "secondary_a": "invoice_ref",
            "secondary_b": "reference",
            "amount_tolerance": "0.05",
            "date_window_days": "2",
        },
    )

    assert response.status_code == 400
    assert "same column for more than one matching role" in response.text


def test_wrong_amount_column_is_rejected_instead_of_silently_scoring_bad_data(
    tmp_path: Path,
) -> None:
    client = make_client(tmp_path)
    prepared = client.post(
        "/prepare",
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": ("payments.csv", B_CSV, "text/csv"),
        },
    )
    stage_match = re.search(r'data-stage-id="([a-f0-9]+)"', prepared.text)
    assert stage_match is not None

    response = client.post(
        "/runs",
        data={
            "stage_id": stage_match.group(1),
            "primary_a": "customer",
            "primary_b": "payer",
            "amount_a": "invoice_ref",
            "amount_b": "reference",
            "amount_tolerance": "0.05",
            "date_window_days": "2",
        },
    )

    assert response.status_code == 400
    assert "does not look like an amount column" in response.text.lower()


def test_invalid_view_uses_html_error_page(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)

    response = client.get(f"{location}?view=not-a-view")

    assert response.status_code == 400
    assert "Request could not be completed" in response.text
    assert "Unknown result view" in response.text
    assert "Start a new reconciliation" in response.text
    assert "Back to start" not in response.text


def test_results_route_paginates_large_runs(tmp_path: Path) -> None:
    from datetime import UTC, datetime

    from rowbridge.models import FieldMapping, MatchDecision, MatchSettings, MatchStatus
    from rowbridge.storage import Repository

    client = make_client(tmp_path)
    repository = Repository(tmp_path / "data" / "rowbridge.sqlite3")
    rows = tuple({"id": f"ITEM-{index:03d}"} for index in range(125))
    decisions = tuple(
        MatchDecision(index, index, 1.0, MatchStatus.AUTO_MATCHED, ())
        for index in range(125)
    )
    run_id = "c" * 32
    repository.save_run(
        run_id=run_id,
        created_at=datetime.now(UTC).isoformat(timespec="seconds"),
        filename_a="a.csv",
        filename_b="b.csv",
        mapping=FieldMapping(primary_a="id", primary_b="id"),
        settings=MatchSettings(),
        rows_a=rows,
        rows_b=rows,
        decisions=decisions,
    )

    page_two = client.get(f"/runs/{run_id}?page=2")

    assert page_two.status_code == 200
    assert "125 rows in this view" in page_two.text
    assert "ITEM-100" in page_two.text
    assert "ITEM-000" not in page_two.text
    assert "Previous" in page_two.text


def test_large_unmatched_sets_switch_manual_link_to_source_row_inputs(tmp_path: Path) -> None:
    from datetime import UTC, datetime

    from rowbridge.models import FieldMapping, MatchDecision, MatchSettings, MatchStatus
    from rowbridge.storage import Repository

    client = make_client(tmp_path)
    repository = Repository(tmp_path / "data" / "rowbridge.sqlite3")
    rows_a = tuple({"id": f"A-{index:03d}"} for index in range(201))
    rows_b = tuple({"id": f"B-{index:03d}"} for index in range(201))
    decisions = tuple(
        [
            MatchDecision(index, None, 0.0, MatchStatus.UNMATCHED, ())
            for index in range(201)
        ]
        + [
            MatchDecision(None, index, 0.0, MatchStatus.UNMATCHED, ())
            for index in range(201)
        ]
    )
    run_id = "d" * 32
    repository.save_run(
        run_id=run_id,
        created_at=datetime.now(UTC).isoformat(timespec="seconds"),
        filename_a="a.csv",
        filename_b="b.csv",
        mapping=FieldMapping(primary_a="id", primary_b="id"),
        settings=MatchSettings(),
        rows_a=rows_a,
        rows_b=rows_b,
        decisions=decisions,
    )

    response = client.get(f"/runs/{run_id}")

    assert response.status_code == 200
    assert 'name="a_row_number"' in response.text
    assert 'name="b_row_number"' in response.text
    assert "This run has many unmatched rows" in response.text
