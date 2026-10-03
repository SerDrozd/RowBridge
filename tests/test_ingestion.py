from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import Workbook

from rowbridge.ingestion import InputFileError, parse_input_bytes, staged_filename


def workbook_bytes(rows: list[list[object]], sheet_name: str = "Data") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_semicolon_csv_detects_delimiter_and_windows_encoding() -> None:
    content = "reference;customer;amount\nA-1;Café Nord;10,00\n".encode("cp1252")

    table = parse_input_bytes(content, "legacy.csv")

    assert table.source_format == "csv"
    assert table.encoding == "cp1252"
    assert table.delimiter == ";"
    assert table.headers == ("reference", "customer", "amount")
    assert table.rows[0]["customer"] == "Café Nord"


def test_utf16_tab_csv_is_supported() -> None:
    content = "id\tname\n1\tAcme\n".encode("utf-16")

    table = parse_input_bytes(content, "utf16.csv")

    assert table.encoding == "utf-16"
    assert table.delimiter == "\t"
    assert table.rows == ({"id": "1", "name": "Acme"},)


def test_xlsx_uses_first_sheet_with_header_and_data() -> None:
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Notes"
    cover.append(["Instructions only"])
    data = workbook.create_sheet("Payments")
    data.append(["reference", "payer", "total"])
    data.append(["INV-1", "Acme Ltd", 10.0])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()

    table = parse_input_bytes(output.getvalue(), "payments.xlsx")

    assert table.source_format == "xlsx"
    assert table.sheet_name == "Payments"
    assert table.headers == ("reference", "payer", "total")
    assert table.rows[0]["total"] == "10"


def test_rejects_unsupported_extension() -> None:
    with pytest.raises(InputFileError, match=r"Use a \.csv or \.xlsx"):
        parse_input_bytes(b"a,b\n1,2\n", "data.txt")


def test_rejects_corrupt_xlsx() -> None:
    with pytest.raises(InputFileError, match="valid XLSX"):
        parse_input_bytes(b"not a workbook", "data.xlsx")


def test_rejects_csv_with_extra_values() -> None:
    with pytest.raises(InputFileError, match="3 values"):
        parse_input_bytes(b"a,b\n1,2,3\n", "broken.csv")


def test_staged_input_path_does_not_use_user_filename(tmp_path: Path) -> None:
    assert staged_filename("a", "../../unsafe.csv") == "a.csv"
    assert staged_filename("b", "strange name.xlsx") == "b.xlsx"
