from __future__ import annotations

import re
from pathlib import Path

from fastapi.testclient import TestClient

from rowbridge.config import Settings
from rowbridge.web import create_app

A_CSV = (
    b"invoice_ref,customer,amount,date\n"
    b"INV-1,Acme Ltd,10.00,2026-10-01\n"
    b"INV-2,Blue Finch GmbH,20.00,2026-10-02\n"
)
B_CSV = (
    b"reference,payer,total,paid_at\n"
    b"INV1,ACME Limited,10.01,2026-10-02\n"
    b"INV-9,Blue Finch GmbH,20.00,2026-10-02\n"
)


def make_client(tmp_path: Path) -> TestClient:
    app = create_app(Settings(data_dir=tmp_path / "data"))
    return TestClient(app)


def create_sample_run(client: TestClient) -> str:
    prepared = client.post(
        "/prepare",
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": ("payments.csv", B_CSV, "text/csv"),
        },
    )
    assert prepared.status_code == 200
    assert "Secondary text" in prepared.text
    stage_match = re.search(r'data-stage-id="([a-f0-9]+)"', prepared.text)
    assert stage_match is not None

    created = client.post(
        "/runs",
        data={
            "stage_id": stage_match.group(1),
            "primary_a": "invoice_ref",
            "primary_b": "reference",
            "secondary_a": "customer",
            "secondary_b": "payer",
            "amount_a": "amount",
            "amount_b": "total",
            "date_a": "date",
            "date_b": "paid_at",
            "amount_tolerance": "0.05",
            "date_window_days": "2",
        },
        follow_redirects=False,
    )
    assert created.status_code == 303
    return created.headers["location"]


def review_match_id(html: str) -> int:
    match = re.search(r'/matches/(\d+)/accept', html)
    assert match is not None
    return int(match.group(1))


def test_vertical_slice_upload_map_persist_render_and_export(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)

    results = client.get(location)
    assert results.status_code == 200
    assert "Reconciliation results" in results.text
    assert "INV-1" in results.text
    assert "INV1" in results.text
    assert "Blue Finch GmbH" in results.text
    assert ">Accept<" in results.text
    assert "Link unmatched rows" in results.text

    fresh_client = make_client(tmp_path)
    persisted = fresh_client.get(location)
    assert persisted.status_code == 200
    assert "INV-1" in persisted.text

    export = fresh_client.get(f"{location}/export.csv")
    assert export.status_code == 200
    assert "status,score,side_a_row" in export.text
    assert "INV-1,INV1" in export.text


def test_accept_review_persists_confirmation_and_audit_event(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)

    accepted = client.post(
        f"{location}/matches/{match_id}/accept",
        follow_redirects=False,
    )
    assert accepted.status_code == 303

    fresh_client = make_client(tmp_path)
    persisted = fresh_client.get(location)
    assert persisted.status_code == 200
    assert "confirmed" in persisted.text
    assert "Accepted proposed match." in persisted.text
    assert "accept" in persisted.text
    assert "<b>0</b><small>Needs review</small>" in persisted.text
    assert "<b>1</b><small>Human matched</small>" in persisted.text

    repeated = fresh_client.post(
        f"{location}/matches/{match_id}/accept",
        follow_redirects=False,
    )
    assert repeated.status_code == 409


def test_reject_review_then_manually_link_unmatched_rows(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)

    rejected = client.post(
        f"{location}/matches/{match_id}/reject",
        follow_redirects=False,
    )
    assert rejected.status_code == 303

    after_reject = client.get(location)
    assert "Rejected proposed match" in after_reject.text
    assert "<b>0</b><small>Needs review</small>" in after_reject.text
    assert "<b>2</b><small>Unmatched</small>" in after_reject.text

    a_option = re.search(r'<option value="(\d+)">row 3 · INV-2</option>', after_reject.text)
    b_option = re.search(r'<option value="(\d+)">row 3 · INV-9</option>', after_reject.text)
    assert a_option is not None
    assert b_option is not None

    linked = client.post(
        f"{location}/manual-links",
        data={
            "a_match_id": a_option.group(1),
            "b_match_id": b_option.group(1),
        },
        follow_redirects=False,
    )
    assert linked.status_code == 303

    fresh_client = make_client(tmp_path)
    persisted = fresh_client.get(location)
    assert "manual matched" in persisted.text
    assert "Linked two previously unmatched rows manually." in persisted.text
    assert "<b>1</b><small>Human matched</small>" in persisted.text
    assert "<b>0</b><small>Unmatched</small>" in persisted.text

    export = fresh_client.get(f"{location}/export.csv")
    assert "manual_matched,,3,3,INV-2,INV-9" in export.text


def test_result_filters_reject_unknown_view(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    location = create_sample_run(client)

    review = client.get(f"{location}?view=review")
    assert review.status_code == 200
    assert "INV-2" in review.text
    assert "INV-1" not in review.text

    invalid = client.get(f"{location}?view=wat")
    assert invalid.status_code == 400



def test_home_file_picker_only_uses_explicit_browse_labels(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.get("/")
    assert response.status_code == 200
    assert 'id="rb-file-a"' in response.text
    assert 'data-file-trigger="rb-file-a"' in response.text
    assert 'id="rb-file-b"' in response.text
    assert 'data-file-trigger="rb-file-b"' in response.text
    assert 'hidden id="rb-file-a"' in response.text
    assert 'hidden id="rb-file-b"' in response.text
    assert 'app.css?v=m4.4' in response.text
    assert 'app.js?v=m4.4' in response.text
    assert 'class="file-input"' not in response.text

def test_prepare_rejects_unsupported_files(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.post(
        "/prepare",
        files={
            "file_a": ("a.txt", A_CSV, "text/plain"),
            "file_b": ("b.csv", B_CSV, "text/csv"),
        },
    )
    assert response.status_code == 400


def xlsx_bytes(headers: list[str], rows: list[list[object]]) -> bytes:
    import io

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def test_prepare_accepts_mixed_csv_and_xlsx_and_shows_preview(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    side_b = xlsx_bytes(
        ["reference", "payer", "total", "paid_at"],
        [["INV1", "ACME Limited", 10.01, "2026-10-02"]],
    )

    response = client.post(
        "/prepare",
        files={
            "file_a": ("orders.csv", A_CSV, "text/csv"),
            "file_b": (
                "payments.xlsx",
                side_b,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )

    assert response.status_code == 200
    assert "Map comparable fields" in response.text
    assert "CSV · utf-8 · comma delimiter" in response.text
    assert "Excel workbook · sheet Data" in response.text
    assert "ACME Limited" in response.text


def test_export_workbook_contains_reconciliation_and_audit_sheets(tmp_path: Path) -> None:
    import io

    from openpyxl import load_workbook

    client = make_client(tmp_path)
    location = create_sample_run(client)
    results = client.get(location)
    match_id = review_match_id(results.text)
    accepted = client.post(f"{location}/matches/{match_id}/accept", follow_redirects=False)
    assert accepted.status_code == 303

    exported = client.get(f"{location}/export.xlsx")
    assert exported.status_code == 200
    assert "spreadsheetml.sheet" in exported.headers["content-type"]

    workbook = load_workbook(io.BytesIO(exported.content), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == [
            "Summary",
            "Reconciliation",
            "Unmatched A",
            "Unmatched B",
            "Review history",
        ]
        reconciliation = workbook["Reconciliation"]
        headers = [cell.value for cell in next(reconciliation.iter_rows())]
        assert headers[:6] == [
            "status",
            "score",
            "side_a_row",
            "side_b_row",
            "side_a_primary",
            "side_b_primary",
        ]
        review = workbook["Review history"]
        review_values = list(review.values)
        assert any(row[1] == "accept" for row in review_values[1:])
    finally:
        workbook.close()


def test_prepare_rejects_unsupported_file_type_with_clear_error(tmp_path: Path) -> None:
    client = make_client(tmp_path)
    response = client.post(
        "/prepare",
        files={
            "file_a": ("a.txt", A_CSV, "text/plain"),
            "file_b": ("b.csv", B_CSV, "text/csv"),
        },
    )
    assert response.status_code == 400
    assert "Use a .csv or .xlsx file" in response.text
