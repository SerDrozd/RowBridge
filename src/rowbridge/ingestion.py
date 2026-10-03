from __future__ import annotations

import csv
import io
import re
import shutil
import time
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from rowbridge.models import CsvTable

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 200
PREVIEW_ROWS = 4
SUPPORTED_SUFFIXES = {".csv", ".xlsx"}
CSV_DELIMITERS = ",;\t|"


class InputFileError(ValueError):
    pass


# Backward-compatible name for older imports while the input layer becomes format-neutral.
CsvInputError = InputFileError


def _validate_size(content: bytes, filename: str) -> None:
    if not content:
        raise InputFileError(f"{filename} is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise InputFileError(f"{filename} is larger than 5 MB")


def _decode_csv(content: bytes, filename: str) -> tuple[str, str]:
    candidates: tuple[str, ...]
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        candidates = ("utf-16",)
    elif content.startswith(b"\xef\xbb\xbf"):
        candidates = ("utf-8-sig",)
    else:
        candidates = ("utf-8", "cp1252")

    for encoding in candidates:
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if "\x00" not in text:
            return text, encoding

    raise InputFileError(
        f"{filename} uses an unsupported text encoding. Save it as UTF-8, UTF-16, "
        "or Windows-1252 and try again."
    )


def _detect_delimiter(text: str) -> str:
    sample = text[:65_536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS)
        return str(dialect.delimiter)
    except csv.Error:
        first_line = next((line for line in text.splitlines() if line.strip()), "")
        counts = {delimiter: first_line.count(delimiter) for delimiter in CSV_DELIMITERS}
        delimiter, count = max(counts.items(), key=lambda item: item[1])
        return delimiter if count else ","


def _normalize_headers(values: list[str], filename: str) -> tuple[str, ...]:
    headers = tuple(value.strip() for value in values)
    if not headers or all(not header for header in headers):
        raise InputFileError(f"{filename} has no header row")
    if len(headers) > MAX_COLUMNS:
        raise InputFileError(f"{filename} exceeds the {MAX_COLUMNS} column limit")
    if any(not header for header in headers):
        raise InputFileError(f"{filename} contains an empty column name")
    if len(set(headers)) != len(headers):
        raise InputFileError(f"{filename} contains duplicate column names")
    return headers


def parse_csv_bytes(content: bytes, filename: str) -> CsvTable:
    _validate_size(content, filename)
    text, encoding = _decode_csv(content, filename)
    delimiter = _detect_delimiter(text)

    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    raw_rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not raw_rows:
        raise InputFileError(f"{filename} has no header row")

    headers = _normalize_headers(raw_rows[0], filename)
    rows: list[dict[str, str]] = []
    for row_number, values in enumerate(raw_rows[1:], start=2):
        if len(rows) >= MAX_ROWS:
            raise InputFileError(f"{filename} exceeds the {MAX_ROWS:,} row limit")
        if len(values) > len(headers):
            raise InputFileError(
                f"{filename} row {row_number} has {len(values)} values but the header has "
                f"{len(headers)} columns"
            )
        padded = values + [""] * (len(headers) - len(values))
        rows.append({header: value.strip() for header, value in zip(headers, padded, strict=True)})

    if not rows:
        raise InputFileError(f"{filename} has no data rows")
    return CsvTable(
        filename=filename,
        headers=headers,
        rows=tuple(rows),
        source_format="csv",
        encoding=encoding,
        delimiter=delimiter,
    )


def _xlsx_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _validate_xlsx_archive(content: bytes, filename: str) -> None:
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            total_size = sum(info.file_size for info in archive.infolist())
            if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                limit_mb = MAX_XLSX_UNCOMPRESSED_BYTES // 1024 // 1024
                raise InputFileError(
                    f"{filename} expands beyond the {limit_mb} MB safety limit"
                )
            names = set(archive.namelist())
            if "xl/workbook.xml" not in names:
                raise InputFileError(f"{filename} is not a valid XLSX workbook")
    except BadZipFile as exc:
        raise InputFileError(f"{filename} is not a valid XLSX workbook") from exc


def parse_xlsx_bytes(content: bytes, filename: str) -> CsvTable:
    _validate_size(content, filename)
    _validate_xlsx_archive(content, filename)

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise InputFileError(f"{filename} could not be read as an XLSX workbook") from exc

    try:
        selected_rows: list[tuple[object, ...]] | None = None
        selected_sheet = ""
        for worksheet in workbook.worksheets:
            non_empty: list[tuple[object, ...]] = []
            for values in worksheet.iter_rows(values_only=True):
                row_values = tuple(values)
                if any(_xlsx_cell_text(value) for value in row_values):
                    non_empty.append(row_values)
                if len(non_empty) > MAX_ROWS + 1:
                    raise InputFileError(f"{filename} exceeds the {MAX_ROWS:,} row limit")
            if len(non_empty) >= 2:
                selected_rows = non_empty
                selected_sheet = worksheet.title
                break

        if selected_rows is None:
            raise InputFileError(f"{filename} has no worksheet with a header and data rows")

        header_values = [_xlsx_cell_text(value) for value in selected_rows[0]]
        while header_values and not header_values[-1]:
            header_values.pop()
        headers = _normalize_headers(header_values, filename)

        rows: list[dict[str, str]] = []
        for values in selected_rows[1:]:
            cell_values = [_xlsx_cell_text(value) for value in values[: len(headers)]]
            padded = cell_values + [""] * (len(headers) - len(cell_values))
            if not any(padded):
                continue
            rows.append(
                {header: value for header, value in zip(headers, padded, strict=True)}
            )
            if len(rows) > MAX_ROWS:
                raise InputFileError(f"{filename} exceeds the {MAX_ROWS:,} row limit")

        if not rows:
            raise InputFileError(f"{filename} has no data rows")
        return CsvTable(
            filename=filename,
            headers=headers,
            rows=tuple(rows),
            source_format="xlsx",
            sheet_name=selected_sheet,
        )
    finally:
        workbook.close()


def parse_input_bytes(content: bytes, filename: str) -> CsvTable:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise InputFileError(f"{filename} is not supported. Use a .csv or .xlsx file.")
    if suffix == ".xlsx":
        return parse_xlsx_bytes(content, filename)
    return parse_csv_bytes(content, filename)


def staged_filename(side: str, original_filename: str) -> str:
    suffix = Path(original_filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise InputFileError(f"{original_filename} is not supported. Use a .csv or .xlsx file.")
    if side not in {"a", "b"}:
        raise ValueError("Staged side must be 'a' or 'b'")
    return f"{side}{suffix}"


def write_staged_input(path: Path, content: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content)


def read_staged_input(path: Path, original_filename: str) -> CsvTable:
    if not path.is_file():
        raise InputFileError("Staged upload was not found. Upload the files again.")
    return parse_input_bytes(path.read_bytes(), original_filename)



_STAGE_ID_PATTERN = re.compile(r"^[0-9a-f]{32}$")


def discard_staged_upload(stage_dir: Path) -> None:
    if stage_dir.is_dir():
        shutil.rmtree(stage_dir, ignore_errors=True)


def cleanup_staged_uploads(
    upload_dir: Path,
    max_age_seconds: int,
    *,
    now: float | None = None,
) -> int:
    if max_age_seconds < 0:
        raise ValueError("max_age_seconds cannot be negative")
    if not upload_dir.is_dir():
        return 0

    cutoff = (time.time() if now is None else now) - max_age_seconds
    removed = 0
    for child in upload_dir.iterdir():
        if not child.is_dir() or not _STAGE_ID_PATTERN.fullmatch(child.name):
            continue
        try:
            modified = child.stat().st_mtime
        except OSError:
            continue
        if modified < cutoff:
            discard_staged_upload(child)
            removed += 1
    return removed


def preview_rows(table: CsvTable) -> tuple[dict[str, str], ...]:
    return table.rows[:PREVIEW_ROWS]
