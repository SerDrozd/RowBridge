from __future__ import annotations

import csv
import io
from collections.abc import Iterable
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from rowbridge.models import MatchStatus, RowPayload
from rowbridge.storage import StoredMatch, StoredReviewEvent, StoredRun


def _source_headers(matches: Iterable[StoredMatch], side: str) -> tuple[str, ...]:
    for match in matches:
        payload = match.a_payload if side == "a" else match.b_payload
        if payload is not None:
            return tuple(payload.keys())
    return ()


def _safe_spreadsheet_text(value: str) -> str:
    stripped = value.lstrip()
    if not stripped:
        return value
    if stripped[0] in {"=", "+", "@", "\t", "\r"}:
        return "'" + value
    if stripped[0] == "-":
        try:
            Decimal(stripped)
        except InvalidOperation:
            return "'" + value
    return value


def _payload_value(payload: RowPayload | None, column: str) -> str:
    if payload is None:
        return ""
    return _safe_spreadsheet_text(payload.get(column, ""))


def _payload_values(payload: RowPayload | None, headers: tuple[str, ...]) -> list[str]:
    return [_payload_value(payload, header) for header in headers]


def _score_text(match: StoredMatch) -> str:
    if match.status == MatchStatus.MANUAL_MATCHED:
        return ""
    return f"{match.score:.4f}"


def _evidence_text(match: StoredMatch) -> str:
    return " | ".join(f"{item.field}: {item.detail}" for item in match.evidence)


def _reconciliation_headers(
    run: StoredRun,
    headers_a: tuple[str, ...],
    headers_b: tuple[str, ...],
) -> list[str]:
    extra_a = tuple(header for header in headers_a if header != run.mapping.primary_a)
    extra_b = tuple(header for header in headers_b if header != run.mapping.primary_b)
    return (
        [
            "status",
            "score",
            "side_a_row",
            "side_b_row",
            "side_a_primary",
            "side_b_primary",
        ]
        + [f"a__{header}" for header in extra_a]
        + [f"b__{header}" for header in extra_b]
        + ["evidence"]
    )


def _reconciliation_rows(
    run: StoredRun,
    matches: tuple[StoredMatch, ...],
    headers_a: tuple[str, ...],
    headers_b: tuple[str, ...],
) -> Iterable[list[object]]:
    extra_a = tuple(header for header in headers_a if header != run.mapping.primary_a)
    extra_b = tuple(header for header in headers_b if header != run.mapping.primary_b)
    for match in matches:
        yield [
            match.status.value,
            _score_text(match),
            match.a_row_number or "",
            match.b_row_number or "",
            _payload_value(match.a_payload, run.mapping.primary_a),
            _payload_value(match.b_payload, run.mapping.primary_b),
            *_payload_values(match.a_payload, extra_a),
            *_payload_values(match.b_payload, extra_b),
            _safe_spreadsheet_text(_evidence_text(match)),
        ]


def build_reconciliation_csv(run: StoredRun, matches: tuple[StoredMatch, ...]) -> str:
    headers_a = _source_headers(matches, "a")
    headers_b = _source_headers(matches, "b")
    output = io.StringIO()
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(_reconciliation_headers(run, headers_a, headers_b))
    writer.writerows(_reconciliation_rows(run, matches, headers_a, headers_b))
    return output.getvalue()


def _append_table(
    worksheet: Any,
    headers: list[str],
    rows: Iterable[list[object]],
) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, column in enumerate(worksheet.columns, start=1):
        width = max((len(str(cell.value or "")) for cell in column), default=0)
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 42)


def _unmatched_rows(
    matches: tuple[StoredMatch, ...], side: str, headers: tuple[str, ...]
) -> Iterable[list[object]]:
    for match in matches:
        if match.status != MatchStatus.UNMATCHED:
            continue
        payload = match.a_payload if side == "a" else match.b_payload
        row_number = match.a_row_number if side == "a" else match.b_row_number
        if payload is not None:
            yield [row_number or "", *_payload_values(payload, headers)]


def _review_rows(events: tuple[StoredReviewEvent, ...]) -> Iterable[list[object]]:
    for event in events:
        yield [
            event.created_at,
            event.action.value,
            event.previous_status or "",
            event.resulting_status,
            event.a_row_number or "",
            event.b_row_number or "",
            "" if event.score is None else f"{event.score:.4f}",
            _safe_spreadsheet_text(event.detail),
        ]


def build_reconciliation_xlsx(
    run: StoredRun,
    matches: tuple[StoredMatch, ...],
    events: tuple[StoredReviewEvent, ...],
) -> bytes:
    headers_a = _source_headers(matches, "a")
    headers_b = _source_headers(matches, "b")
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Summary"
    summary_rows = [
        ("Run ID", run.id),
        ("Created", run.created_at),
        ("Side A", run.filename_a),
        ("Side B", run.filename_b),
        ("Side A rows", run.total_a),
        ("Side B rows", run.total_b),
        ("Primary mapping", f"{run.mapping.primary_a} <-> {run.mapping.primary_b}"),
    ]
    for label, value in summary_rows:
        summary.append([label, value])
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 52

    reconciliation = workbook.create_sheet("Reconciliation")
    _append_table(
        reconciliation,
        _reconciliation_headers(run, headers_a, headers_b),
        _reconciliation_rows(run, matches, headers_a, headers_b),
    )

    unmatched_a = workbook.create_sheet("Unmatched A")
    _append_table(
        unmatched_a,
        ["source_row", *headers_a],
        _unmatched_rows(matches, "a", headers_a),
    )

    unmatched_b = workbook.create_sheet("Unmatched B")
    _append_table(
        unmatched_b,
        ["source_row", *headers_b],
        _unmatched_rows(matches, "b", headers_b),
    )

    review = workbook.create_sheet("Review history")
    _append_table(
        review,
        [
            "created_at",
            "action",
            "previous_status",
            "resulting_status",
            "side_a_row",
            "side_b_row",
            "score",
            "detail",
        ],
        _review_rows(events),
    )

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
